# -*- coding: utf-8 -*-
"""
自动生成 Word 分析报告（含论文规范的三线表与图表）。
三线表 = 顶线 + 表头底线 + 底线，无竖线、无行间横线。
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def _set_three_line_borders(table) -> None:
    """把普通表格改成三线表样式。"""
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge, size in (("top", 12), ("bottom", 12)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), "000000")
        borders.append(el)
    for edge in ("left", "right", "insideV", "insideH"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tbl_pr.append(borders)

    # 表头行底线
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_borders = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "000000")
        tc_borders.append(bottom)
        tc_pr.append(tc_borders)


def _add_table(doc: Document, df: pd.DataFrame, font_size: int = 9) -> None:
    """把 DataFrame 渲染为三线表。"""
    table = doc.add_table(rows=df.shape[0] + 1, cols=df.shape[1])
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for j, col in enumerate(df.columns):
        cell = table.rows[0].cells[j]
        cell.text = str(col)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(font_size)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    for i, (_, row) in enumerate(df.iterrows(), start=1):
        for j, value in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = str(value)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(font_size)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    _set_three_line_borders(table)
    doc.add_paragraph()


def _set_east_asia(run, font_name: str) -> None:
    """为 run 设置中文字体（rPr/rFonts 不存在时自动创建）。"""
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), font_name)


def _set_default_font(doc: Document) -> None:
    """正文宋体五号、西文 Times New Roman。"""
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)
    rpr = style.element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")


def _heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1f, 0x3d, 0x2e)
        _set_east_asia(run, "黑体")


def build_report(context: dict, out_dir: Path) -> Path:
    """根据 main.py 传入的分析结果上下文生成 docx 报告。"""
    doc = Document()
    _set_default_font(doc)

    # ---- 封面信息 ----
    title = doc.add_heading("问卷数据自动分析报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        _set_east_asia(run, "黑体")

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}\n"
                 "生成方式：survey-analysis-pipeline 全自动分析管线\n"
                 "数据说明：仿真问卷数据（N=320），仅用于管线演示，不代表真实结论")

    # ---- 一、样本清洗与质量控制 ----
    _heading(doc, "一、样本清洗与质量控制")
    flow = pd.DataFrame(list(context["clean_report"].items()),
                        columns=["环节", "数值"])
    _add_table(doc, flow)
    doc.add_paragraph(
        "清洗规则：量表越界值置空并以题项中位数插补；量表题缺失比例超过 20% 的问卷整份剔除；"
        "直线作答（8 道量表题答案完全一致）与注意力检验题作答错误的问卷不进入分析样本。")

    # ---- 二、信度分析 ----
    _heading(doc, "二、量表信度分析（Cronbach's α）")
    _add_table(doc, context["reliability"])

    # ---- 三、描述统计 ----
    _heading(doc, "三、描述统计")
    _heading(doc, "3.1 分类变量频数分布（示例：学历）", level=2)
    _add_table(doc, context["freq_edu"])
    _heading(doc, "3.2 数值变量与量表得分描述统计", level=2)
    _add_table(doc, context["describe"])
    if context.get("fig_scale_box"):
        doc.add_picture(str(context["fig_scale_box"]), width=Cm(15))
        cap = doc.add_paragraph("图 1　各量表维度得分箱线图")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ---- 四、交叉分析与卡方检验 ----
    _heading(doc, "四、交叉分析与卡方检验")
    cross = context["crosstab"]
    _heading(doc, "4.1 学历 × 购买情况 列联表", level=2)
    _add_table(doc, cross["table"])
    _heading(doc, "4.2 卡方检验结果", level=2)
    _add_table(doc, cross["result"])
    p = doc.add_paragraph()
    p.add_run("自动结论：").bold = True
    p.add_run(cross["conclusion"])
    if context.get("fig_purchase_rate"):
        doc.add_picture(str(context["fig_purchase_rate"]), width=Cm(15))
        cap = doc.add_paragraph("图 2　不同学历组购买率")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ---- 五、Logistic 回归 ----
    _heading(doc, "五、购买行为影响因素：二元 Logistic 回归")
    _add_table(doc, context["regression_table"])
    doc.add_paragraph(
        f"模型概况：N = {context['regression_meta']['样本量N']}，"
        f"McFadden Pseudo R² = {context['regression_meta']['Pseudo R²']}，"
        f"整体模型 p = {context['regression_meta']['模型p值']}。"
        "显著性标记：* p<0.05，** p<0.01，*** p<0.001。"
        "OR>1 表示该变量取值越高，购买发生比越大。")
    if context.get("fig_corr"):
        doc.add_picture(str(context["fig_corr"]), width=Cm(13))
        cap = doc.add_paragraph("图 3　核心变量相关性热力图")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ---- 六、开放题文本 ----
    _heading(doc, "六、开放题文本关键词")
    _add_table(doc, context["keywords"])
    if context.get("fig_keywords"):
        doc.add_picture(str(context["fig_keywords"]), width=Cm(14))
        cap = doc.add_paragraph("图 4　开放题高频关键词")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    path = out_dir / "问卷数据分析报告.docx"
    doc.save(str(path))
    return path
